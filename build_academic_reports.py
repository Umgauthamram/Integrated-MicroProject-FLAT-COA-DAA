"""
build_academic_reports.py - Generates 3 publication-quality academic microproject reports
for FLAT, COA, and DAA in exact accordance with the Alliance University / Kalvium format
template ('Principles of Science Microproject - Wi-Fi Signal Attenuation.docx').
Produces both native Word (.docx) and official PDF (.pdf) documents.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client

NAVY = RGBColor(0x1A, 0x36, 0x5D)
BLUE = RGBColor(0x2B, 0x6C, 0xB0)
DARK_GRAY = RGBColor(0x33, 0x41, 0x55)
LIGHT_BG = "F1F5F9"
BORDER_COLOR = "CBD5E1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report_document(subject_title, subject_code_name, project_title, sections_data, table_data, screenshot_path, mentor_name="Faculty Mentor"):
    doc = docx.Document()

    # Set 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = DARK_GRAY

    # ==================== COVER PAGE ====================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("ALLIANCE UNIVERSITY\n")
    r_inst.font.size = Pt(18)
    r_inst.font.bold = True
    r_inst.font.color.rgb = NAVY

    r_prog = p_inst.add_run("Kalvium Undergraduate Programme\n")
    r_prog.font.size = Pt(14)
    r_prog.font.color.rgb = BLUE

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(f"{subject_title.upper()}\n")
    r_sub.font.size = Pt(16)
    r_sub.font.bold = True
    r_sub.font.color.rgb = NAVY

    r_micro = p_sub.add_run("MICROPROJECT REPORT\n\n")
    r_micro.font.size = Pt(14)
    r_micro.font.color.rgb = DARK_GRAY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"\"{project_title}\"\n")
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    # University Seal Image
    base_dir = os.path.dirname(os.path.abspath(__file__))
    seal_path = os.path.join(base_dir, 'word', 'media', 'image1.png')
    template_docx = os.path.join(base_dir, 'Principles of Science Microproject - Wi-Fi Signal Attenuation.docx')

    if not os.path.exists(seal_path) and os.path.exists(template_docx):
        import zipfile
        with zipfile.ZipFile(template_docx, 'r') as z:
            z.extract('word/media/image1.png', base_dir)
    if os.path.exists(seal_path):
        p_seal = doc.add_paragraph()
        p_seal.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_seal.paragraph_format.space_before = Pt(10)
        p_seal.paragraph_format.space_after = Pt(20)
        doc.add_picture(seal_path, width=Inches(2.1))
        p_seal_img = doc.paragraphs[-1]
        p_seal_img.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Submission Metadata
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(15)
    p_meta.paragraph_format.line_spacing = 1.3

    meta_items = [
        ("Submitted by", "Gautham Ram"),
        ("Programme", "B.Tech. (Computer Science & Engineering)"),
        ("Institution", "Kalvium | Alliance University"),
        ("Squad", "70"),
        ("Course Faculty / Mentor", mentor_name),
        ("Submission Date", "October 2026")
    ]
    for label, val in meta_items:
        r_lbl = p_meta.add_run(f"{label}: ")
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(11.5)
        r_lbl.font.color.rgb = NAVY
        r_val = p_meta.add_run(f"{val}\n")
        r_val.font.size = Pt(11.5)
        r_val.font.color.rgb = DARK_GRAY

    doc.add_page_break()

    # ==================== NUMBERED SECTIONS ====================
    for sec_num, sec_title, sec_paragraphs in sections_data:
        # Section Heading
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r_h = h.add_run(f"{sec_num}. {sec_title}")
        r_h.font.size = Pt(15)
        r_h.font.bold = True
        r_h.font.color.rgb = NAVY

        # Paragraphs
        for p_item in sec_paragraphs:
            if isinstance(p_item, tuple):
                p_type, p_text = p_item
            else:
                p_type, p_text = "body", p_item

            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)

            if p_type == "formula":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run(p_text)
                r.font.name = 'Consolas'
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = NAVY
            elif p_type == "bullet":
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(4)
                r_bullet = p.add_run("•  ")
                r_bullet.font.bold = True
                r_bullet.font.color.rgb = BLUE
                r = p.add_run(p_text)
                r.font.size = Pt(10.5)
            elif p_type == "subheading":
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(p_text)
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = BLUE
            else:
                r = p.add_run(p_text)
                r.font.size = Pt(11)

        # In Section 7, insert the Observational / Benchmark Table & Screenshot
        if sec_num == 7:
            # 1. Insert Data Table
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            tbl_title = doc.add_paragraph()
            tbl_title.paragraph_format.space_before = Pt(8)
            tbl_title.paragraph_format.space_after = Pt(4)
            tbl_title.paragraph_format.keep_with_next = True
            r_tt = tbl_title.add_run(f"Table 7.1: {subject_code_name} Experimental Verification & Benchmark Metrics")
            r_tt.font.size = Pt(11)
            r_tt.font.bold = True
            r_tt.font.color.rgb = BLUE

            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            for row_idx, row_cells in enumerate(table_data):
                row = table.rows[row_idx]
                is_header = (row_idx == 0)
                for col_idx, cell_value in enumerate(row_cells):
                    cell = row.cells[col_idx]
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    if is_header:
                        set_cell_background(cell, "1A365D")
                    else:
                        set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")

                    cp = cell.paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_header or col_idx > 0) else WD_ALIGN_PARAGRAPH.LEFT
                    cp.paragraph_format.space_before = Pt(3)
                    cp.paragraph_format.space_after = Pt(3)
                    cr = cp.add_run(str(cell_value))
                    cr.font.name = 'Consolas' if (not is_header and col_idx >= 1) else 'Calibri'
                    cr.font.size = Pt(9.5) if not is_header else Pt(10)
                    if is_header:
                        cr.font.bold = True
                        cr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        cr.font.color.rgb = DARK_GRAY

            # 2. Insert Screenshot
            if screenshot_path and os.path.exists(screenshot_path):
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                p_fig = doc.add_paragraph()
                p_fig.paragraph_format.space_before = Pt(10)
                p_fig.paragraph_format.space_after = Pt(4)
                p_fig.paragraph_format.keep_with_next = True
                p_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_fig = p_fig.add_run(f"Figure 7.1: {subject_code_name} Interactive Module Simulation Output & Diagnostic View")
                r_fig.font.size = Pt(11)
                r_fig.font.bold = True
                r_fig.font.color.rgb = BLUE

                doc.add_picture(screenshot_path, width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc

def convert_docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        print(f"  -> Converted to PDF: {pdf_path}")
    finally:
        word.Quit()
